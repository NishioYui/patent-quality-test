from __future__ import annotations

from pathlib import Path
import json

from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUTPUT_DOCX_NAME = "patent_draft.docx"


def _read_json(path: Path) -> dict:
    if not path.exists():
        return {}
    try:
        obj = json.loads(path.read_text(encoding="utf-8"))
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}


def _set_japanese_font(run, font_name: str = "Meiryo") -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(22)
    section.right_margin = Mm(22)

    normal = doc.styles["Normal"]
    normal.font.name = "Meiryo"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Meiryo")
    normal.font.size = Pt(10.5)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Meiryo"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Meiryo")
    h1.font.size = Pt(13)
    h1.font.bold = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Meiryo"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Meiryo")
    h2.font.size = Pt(11.5)
    h2.font.bold = True


def _add_title(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(title)
    _set_japanese_font(r)
    r.bold = True
    r.font.size = Pt(15)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def _add_text_block(doc: Document, text: str) -> None:
    if not text:
        return
    for para in str(text).split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(para)
        _set_japanese_font(r)


def _add_numbered_claim(doc: Document, claim_no: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25

    r1 = p.add_run(f"【請求項{claim_no}】")
    _set_japanese_font(r1)
    r1.bold = True

    r2 = p.add_run(text or "")
    _set_japanese_font(r2)


def _write_claims(doc: Document, claims_obj: dict) -> None:
    claims = claims_obj.get("claims", {}) if isinstance(claims_obj.get("claims"), dict) else {}
    independent = claims.get("independent", {}) if isinstance(claims.get("independent"), dict) else {}
    dependent = claims.get("dependent", []) if isinstance(claims.get("dependent"), list) else []

    _add_heading(doc, "特許請求の範囲", level=1)

    if independent:
        _add_numbered_claim(doc, 1, independent.get("text", ""))

    for item in dependent:
        if not isinstance(item, dict):
            continue
        cid = str(item.get("id", ""))
        claim_no = 0
        if cid.startswith("D") and cid[1:].isdigit():
            claim_no = int(cid[1:])
        if claim_no <= 0:
            continue
        _add_numbered_claim(doc, claim_no, item.get("text", ""))


def _write_patent_documents(doc: Document, docs: list[dict]) -> None:
    _add_heading(doc, "先行技術文献", level=2)

    if not docs:
        _add_text_block(doc, "なし")
        return

    for item in docs:
        if not isinstance(item, dict):
            continue
        label = item.get("label", "")
        citation = item.get("citation", "")
        line = f"{label}　{citation}".strip()
        _add_text_block(doc, line)


def _write_drawings(doc: Document, drawings: list[dict]) -> None:
    _add_heading(doc, "図面の簡単な説明", level=2)

    if not drawings:
        _add_text_block(doc, "図面なし。")
        return

    for item in drawings:
        if not isinstance(item, dict):
            continue
        label = item.get("label", "")
        desc = item.get("description", "")
        line = f"{label}：{desc}".strip("：")
        _add_text_block(doc, line)


def _write_mode_for_carrying_out(doc: Document, mode: dict) -> None:
    _add_heading(doc, "発明を実施するための形態", level=2)

    fields = [
        ("【概要】", mode.get("preface", "")),
        ("【構成】", mode.get("configuration", "")),
        ("【動作】", mode.get("operation", "")),
        ("【効果】", mode.get("effect", "")),
        ("【変形例】", mode.get("modifications_note", "")),
    ]

    for label, value in fields:
        if not value:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label)
        _set_japanese_font(r)
        r.bold = True
        _add_text_block(doc, value)


def _write_specification(doc: Document, spec_obj: dict) -> None:
    spec = spec_obj.get("specification", {}) if isinstance(spec_obj.get("specification"), dict) else {}

    _add_heading(doc, "明細書", level=1)

    _add_heading(doc, "発明の名称", level=2)
    _add_text_block(doc, spec.get("invention_title", ""))

    _add_heading(doc, "技術分野", level=2)
    _add_text_block(doc, spec.get("technical_field", ""))

    _add_heading(doc, "背景技術", level=2)
    _add_text_block(doc, spec.get("background_art", ""))

    _add_heading(doc, "従来技術文献", level=2)
    _add_text_block(doc, spec.get("prior_art_literature", ""))

    _write_patent_documents(doc, spec.get("patent_documents", []))

    _add_heading(doc, "発明の概要", level=2)
    _add_text_block(doc, spec.get("summary_of_invention", ""))

    _add_heading(doc, "発明が解決しようとする課題", level=2)
    _add_text_block(doc, spec.get("problem_to_be_solved", ""))

    _add_heading(doc, "課題を解決するための手段", level=2)
    _add_text_block(doc, spec.get("means_for_solving", ""))

    _add_heading(doc, "発明の効果", level=2)
    _add_text_block(doc, spec.get("effect_of_invention", ""))

    _write_drawings(doc, spec.get("drawings", []))
    _write_mode_for_carrying_out(doc, spec.get("mode_for_carrying_out", {}))

    _add_heading(doc, "符号の説明", level=2)
    _add_text_block(doc, spec.get("reference_signs", ""))


def _write_warnings(doc: Document, claims_obj: dict, spec_obj: dict) -> None:
    warnings = []
    for obj in (claims_obj, spec_obj):
        ws = obj.get("warnings", [])
        if isinstance(ws, list):
            warnings.extend(ws)

    if not warnings:
        return

    _add_heading(doc, "警告一覧", level=1)

    for idx, w in enumerate(warnings, start=1):
        if not isinstance(w, dict):
            continue
        code = str(w.get("code", ""))
        severity = str(w.get("severity", ""))
        message = str(w.get("message", ""))

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2

        head = f"{idx}. [{severity} / {code}] "
        r1 = p.add_run(head)
        _set_japanese_font(r1)
        r1.bold = True

        r2 = p.add_run(message)
        _set_japanese_font(r2)

        evidence = w.get("evidence", [])
        if isinstance(evidence, list) and evidence:
            for ev in evidence:
                if not isinstance(ev, dict):
                    continue
                quote = ev.get("quote", "")
                location = ev.get("location", "")
                _add_text_block(doc, f"  - {location}: {quote}")


def build_patent_docx(jdir: Path, include_warnings: bool = True) -> Path:
    claims_path = jdir / "claims.json"
    spec_path = jdir / "spec.json"

    claims_obj = _read_json(claims_path)
    spec_obj = _read_json(spec_path)

    if not claims_obj and not spec_obj:
        raise FileNotFoundError("claims.json and spec.json are both missing or unreadable")

    spec = spec_obj.get("specification", {}) if isinstance(spec_obj.get("specification"), dict) else {}
    invention_title = spec.get("invention_title", "") or "特許出願書類案"

    doc = Document()
    _configure_document(doc)

    _add_title(doc, invention_title)
    _write_claims(doc, claims_obj)
    _write_specification(doc, spec_obj)

    if include_warnings:
        _write_warnings(doc, claims_obj, spec_obj)

    out_path = jdir / OUTPUT_DOCX_NAME
    doc.save(str(out_path))
    return out_path